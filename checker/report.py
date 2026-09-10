"""Экспорт отчёта в .xlsx (раскрашенные листы) и .csv.

Исходный файл никогда не изменяется — формируется отдельный отчёт в память.
"""
from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Issue, Level

_FILLS = {
    Level.ERROR: PatternFill("solid", fgColor="F8CBAD"),
    Level.WARNING: PatternFill("solid", fgColor="FFE699"),
    Level.ADVICE: PatternFill("solid", fgColor="BDD7EE"),
    Level.TECH: PatternFill("solid", fgColor="D9D9D9"),
}
_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

_COLUMNS = ["Уровень", "Лист", "Строка", "Колонка", "Бренд", "Тип поста",
            "Статус", "Суть замечания", "Как исправить", "Код", "Ссылка"]


def _issue_row(iss: Issue) -> list[Any]:
    return [
        f"{iss.level.emoji} {iss.level.title_ru}",
        iss.sheet, iss.row, iss.column, iss.brand, iss.post_type,
        iss.status, iss.message, iss.fix, iss.code, iss.link,
    ]


def _write_sheet(ws, issues: list[Issue]) -> None:
    ws.append(_COLUMNS)
    for c in range(1, len(_COLUMNS) + 1):
        cell = ws.cell(1, c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for iss in issues:
        ws.append(_issue_row(iss))
        fill = _FILLS.get(iss.level)
        if fill:
            for c in range(1, len(_COLUMNS) + 1):
                ws.cell(ws.max_row, c).fill = fill
    widths = [18, 16, 8, 16, 8, 16, 12, 55, 55, 20, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLUMNS))}{ws.max_row}"
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 8).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 9).alignment = Alignment(wrap_text=True, vertical="top")


def build_xlsx(issues: list[Issue], summary: dict[str, Any],
               holiday_issues: list[Issue] | None = None,
               online_issues: list[Issue] | None = None) -> bytes:
    wb = Workbook()

    # Сводка
    ws = wb.active
    ws.title = "Сводка"
    ws.append(["Показатель", "Значение"])
    for c in (1, 2):
        ws.cell(1, c).fill = _HEADER_FILL
        ws.cell(1, c).font = _HEADER_FONT
    for k, v in summary.items():
        ws.append([k, v])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 20

    reg = [i for i in issues if i.sheet == "Реестр постов"]
    txt = [i for i in issues if i.sheet != "Реестр постов"]

    _write_sheet(wb.create_sheet("Реестр"), reg)
    _write_sheet(wb.create_sheet("Тексты"), txt)
    if holiday_issues is not None:
        _write_sheet(wb.create_sheet("Праздники"), holiday_issues)
    if online_issues is not None:
        _write_sheet(wb.create_sheet("Онлайн"), online_issues)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(issues: list[Issue]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(_COLUMNS)
    for iss in issues:
        writer.writerow([
            f"{iss.level.emoji} {iss.level.title_ru}", iss.sheet, iss.row,
            iss.column, iss.brand, iss.post_type, iss.status,
            iss.message, iss.fix, iss.code, iss.link,
        ])
    return buf.getvalue().encode("utf-8-sig")
