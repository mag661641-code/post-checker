"""Чтение Excel-файла: объединённые ячейки, поиск колонок по названию,
группировка постов на листах брендов.

Читаем через openpyxl, чтобы корректно разворачивать объединённые ячейки.
Модуль не зависит от Streamlit.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any, Optional

import openpyxl

from . import normalize as N
from .models import PostRecord, RegistryRow

REGISTRY_SHEET = "Реестр постов"
BRAND_SHEETS = ["СМУ", "ИМП", "МПЭ", "МПИ", "АПС"]
HOLIDAYS_SHEET = "Обязательные праздники"
SHIPMENTS_SHEET = "Отгрузки"

# как может называться колонка с датой поста на листах брендов (по приоритету)
DATE_COL_NAMES = ["дата", "когда выложить", "дата публикации (план)",
                  "дата публикации"]

_DATE_IN_TEXT_RE = re.compile(r"\d{1,2}\.\d{1,2}\.\d{4}")


def _date_and_note(raw_val: Any) -> tuple[Optional[dt.date], str]:
    """Дата и примечание из ячейки.

    Если значение — обычная дата, вернуть (дата, ""). Если это текст с датой
    и примечанием («18.07.2025 (доп)», «30.06.2025\nЯБ 01.07.2025»), взять
    первую дату по шаблону, а всю исходную запись сохранить как примечание.
    """
    d, _ = N.parse_date(raw_val)
    if d is not None:
        return d, ""
    if raw_val is None:
        return None, ""
    s = re.sub(r"\s+", " ", str(raw_val)).strip()
    if not s:
        return None, ""
    m = _DATE_IN_TEXT_RE.search(s)
    if m:
        d2, _ = N.parse_date(m.group(0))
        if d2 is not None:
            return d2, f"в таблице указано: {s}"
    return None, ""


@dataclass
class LoadedWorkbook:
    """Результат чтения файла."""
    sheet_names: list[str] = field(default_factory=list)
    registry: list[RegistryRow] = field(default_factory=list)
    posts: dict[str, list[PostRecord]] = field(default_factory=dict)   # по бренду-листу
    stats_comments: list[dict] = field(default_factory=list)           # блок «Комментарии»
    holidays_raw: list[dict] = field(default_factory=list)
    sheet_info: dict[str, dict] = field(default_factory=dict)          # для страницы «Загрузка»
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Разворачивание объединённых ячеек
# ---------------------------------------------------------------------------
def _build_grid(ws) -> list[list[Any]]:
    """Прочитать лист в матрицу значений, развернув объединённые ячейки."""
    return _build_grids(ws)[1]


def _build_grids(ws) -> tuple[list[list[Any]], list[list[Any]]]:
    """Вернуть (raw_grid, expanded_grid).

    raw_grid — как в файле (в объединённой ячейке заполнена только верхняя левая).
    expanded_grid — значение объединённой ячейки развёрнуто на весь диапазон.
    Raw нужен, чтобы находить начало группы поста (текст в «Пост» есть только
    в верхней строке объединённого блока).
    """
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    raw = [[ws.cell(r + 1, c + 1).value for c in range(max_col)]
           for r in range(max_row)]
    expanded = [list(row) for row in raw]
    for mr in ws.merged_cells.ranges:
        top = expanded[mr.min_row - 1][mr.min_col - 1]
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                expanded[r - 1][c - 1] = top
    return raw, expanded


def _find_header_row(grid: list[list[Any]], required: list[str],
                     scan: int = 6) -> Optional[int]:
    """Найти строку заголовка по наличию нужных названий колонок."""
    req = [r.strip().lower() for r in required]
    for i in range(min(scan, len(grid))):
        cells = [str(v).strip().lower() if v is not None else "" for v in grid[i]]
        hits = sum(1 for r in req if any(r == c or (r and r in c) for c in cells))
        if hits >= max(2, len(req) // 2):
            return i
    return None


def _column_map(header: list[Any]) -> dict[str, int]:
    """Сопоставить нормализованное имя колонки -> индекс (0-based)."""
    m: dict[str, int] = {}
    for idx, v in enumerate(header):
        if v is None:
            continue
        key = re.sub(r"\s+", " ", str(v)).strip().lower()
        if key and key not in m:
            m[key] = idx
    return m


def _get(colmap: dict[str, int], row: list[Any], *names: str) -> Any:
    for name in names:
        idx = colmap.get(name.strip().lower())
        if idx is not None and idx < len(row):
            return row[idx]
    return None


# ---------------------------------------------------------------------------
# Главный вход
# ---------------------------------------------------------------------------
def load_workbook(source) -> LoadedWorkbook:
    """source — путь к файлу, bytes или файловый объект Streamlit."""
    if isinstance(source, (bytes, bytearray)):
        source = BytesIO(source)
    wb = openpyxl.load_workbook(source, data_only=True)
    result = LoadedWorkbook(sheet_names=list(wb.sheetnames))

    for name in wb.sheetnames:
        ws = wb[name]
        result.sheet_info[name] = {"rows": ws.max_row, "cols": ws.max_column}

    if REGISTRY_SHEET in wb.sheetnames:
        try:
            result.registry = _load_registry(wb[REGISTRY_SHEET])
            result.sheet_info[REGISTRY_SHEET]["parsed"] = len(result.registry)
        except Exception as e:  # noqa: BLE001
            result.warnings.append(f"Лист «{REGISTRY_SHEET}»: {e}")

    for name in BRAND_SHEETS:
        if name in wb.sheetnames:
            try:
                posts, comments = _load_brand_sheet(wb[name], name)
                result.posts[name] = posts
                result.stats_comments.extend(comments)
                result.sheet_info[name]["parsed"] = len(posts)
            except Exception as e:  # noqa: BLE001
                result.warnings.append(f"Лист «{name}»: {e}")
                result.posts[name] = []

    if HOLIDAYS_SHEET in wb.sheetnames:
        try:
            result.holidays_raw = _load_holidays(wb[HOLIDAYS_SHEET])
        except Exception as e:  # noqa: BLE001
            result.warnings.append(f"Лист «{HOLIDAYS_SHEET}»: {e}")

    return result


# ---------------------------------------------------------------------------
# Реестр постов
# ---------------------------------------------------------------------------
def _load_registry(ws) -> list[RegistryRow]:
    grid = _build_grid(ws)
    hdr_i = _find_header_row(grid, ["бренд", "тип поста", "ссылка", "статус"])
    if hdr_i is None:
        raise ValueError("не найдена строка заголовка")
    colmap = _column_map(grid[hdr_i])
    rows: list[RegistryRow] = []
    for i in range(hdr_i + 1, len(grid)):
        row = grid[i]
        excel_row = i + 1
        brand = _clean(_get(colmap, row, "бренд"))
        links_raw = _get(colmap, row, "ссылка")
        status = _clean(_get(colmap, row, "статус"))
        ptype = _clean(_get(colmap, row, "тип поста", "тип"))
        # пустая строка целиком — пропускаем
        if not any([brand, status, ptype, links_raw]):
            continue
        rr = RegistryRow(
            row=excel_row,
            write_date_raw=_get(colmap, row, "дата написания"),
            pub_date_raw=_get(colmap, row, "дата публикации (план)",
                              "дата публикации"),
            brand=brand,
            executor=_clean(_get(colmap, row, "исполнитель")),
            post_type=ptype,
            links_raw=str(links_raw) if links_raw is not None else "",
            status=status,
        )
        rr.write_date, _ = N.parse_date(rr.write_date_raw)
        rr.pub_date, _ = N.parse_date(rr.pub_date_raw)
        rr.links = N.extract_links(rr.links_raw)
        rr.links_norm = [N.normalize_link(u) for u in rr.links]
        rows.append(rr)
    return rows


# ---------------------------------------------------------------------------
# Листы брендов
# ---------------------------------------------------------------------------
def _load_brand_sheet(ws, brand: str) -> tuple[list[PostRecord], list[dict]]:
    raw, grid = _build_grids(ws)
    hdr_i = _find_header_row(grid, ["соцсеть", "ссылка", "пост", "фото"])
    if hdr_i is None:
        raise ValueError("не найдена строка заголовка")
    colmap = _column_map(grid[hdr_i])
    post_idx = colmap.get("пост")

    # колонка даты: у разных листов называется по-разному (на СМУ — «Когда выложить»)
    date_key = next((n for n in DATE_COL_NAMES if n in colmap), "дата")
    date_col = colmap.get(date_key)
    soc_col = colmap.get("соцсеть")
    link_col = colmap.get("ссылка")

    posts: list[PostRecord] = []
    comments: list[dict] = []
    current: Optional[PostRecord] = None

    def flush():
        nonlocal current
        if current is not None:
            posts.append(current)
        current = None

    for i in range(hdr_i + 1, len(grid)):
        row = grid[i]
        excel_row = i + 1
        ptype = _clean(_get(colmap, row, "тип"))
        photo = _clean(_get(colmap, row, "фото"))
        stat = _get(colmap, row, "статистика")
        executor = _clean(_get(colmap, row, "исполнитель"))

        # дату, текст, соцсеть и ссылку берём из «сырой» сетки — чтобы значение
        # объединённой ячейки не «протекало» в строки-продолжения
        raw_date = raw[i][date_col] if date_col is not None and i < len(raw) else None
        raw_text = raw[i][post_idx] if post_idx is not None and i < len(raw) else None
        social = _clean(raw[i][soc_col]) if soc_col is not None and i < len(raw) else ""
        link = _clean(raw[i][link_col]) if link_col is not None and i < len(raw) else ""
        has_text = bool(raw_text and str(raw_text).strip())

        # служебная строка с годом («2025», «2026») или месяцем («Апрель 2026»)
        if _is_service_row(row, colmap, raw_date, social, link, raw_text):
            continue

        parsed_date, date_note = _date_and_note(raw_date)

        # Новый пост начинается на строке с датой; либо на строке с текстом
        # и соцсетью, но без даты (это настоящий пост без даты).
        starts_new = parsed_date is not None or (has_text and bool(social))

        if starts_new:
            flush()
            current = PostRecord(
                sheet=brand, row=excel_row, brand=brand,
                date=parsed_date, date_note=date_note, post_type=ptype,
                text=str(raw_text) if has_text else "",
                executor=executor, raw_rows=[excel_row],
            )
            link_in_text = _cell_hyperlink(ws, excel_row, post_idx)
            if link_in_text:
                current.text_links.append(link_in_text)
            if photo:
                current.photos.append(photo)
            if social or link:
                current.socials.append({"social": social, "link": link, "row": excel_row})
            _collect_stat_comment(stat, brand, excel_row, comments, current)
        else:
            if current is None:
                # строка-продолжение без начала — пропускаем
                continue
            current.raw_rows.append(excel_row)
            # текст на строке-продолжении (напр. «Подпись:», «Клиент:», версия
            # для другой соцсети) присоединяем к посту и проверяем вместе
            if has_text:
                current.text = (current.text + "\n\n" + str(raw_text)).strip()
                link_in_text = _cell_hyperlink(ws, excel_row, post_idx)
                if link_in_text and link_in_text not in current.text_links:
                    current.text_links.append(link_in_text)
            if social or link:
                current.socials.append({"social": social, "link": link, "row": excel_row})
            if photo:
                current.photos.append(photo)
            _collect_stat_comment(stat, brand, excel_row, comments, current)

    flush()
    return posts, comments


def _cell_hyperlink(ws, row: int, col_idx: Optional[int]) -> Optional[str]:
    """Ссылка, вшитая в ячейку целиком (Insert → Link). Частичные анкоры
    и формулы HYPERLINK через openpyxl недоступны — их читаем из Google API."""
    if col_idx is None:
        return None
    try:
        cell = ws.cell(row=row, column=col_idx + 1)
        h = getattr(cell, "hyperlink", None)
        target = getattr(h, "target", None) if h else None
        return target or None
    except Exception:  # noqa: BLE001
        return None


def _collect_stat_comment(stat: Any, brand: str, row: int,
                          comments: list[dict], post: PostRecord) -> None:
    """Комментарии руководителя из колонки «Статистика» (не число)."""
    if stat is None:
        return
    s = str(stat).strip()
    if not s:
        return
    # числа и «доли» — это статистика, а не комментарий
    if re.fullmatch(r"[\d\s.,%()+\-–—]+", s):
        return
    # заголовки служебных таблиц
    if s.lower() in {"статистика", "вконтакте", "одноклассники", "план", "факт"}:
        return
    comments.append({"brand": brand, "row": row, "text": s})
    post.stats_comments.append(s)


def _is_service_row(row, colmap, date_val, social, link, text) -> bool:
    """Служебная строка: только год/месяц без данных поста."""
    if social or link or text:
        return False
    if date_val is None:
        return True
    s = str(date_val).strip()
    if re.fullmatch(r"20\d\d(\.0)?", s):
        return True
    if re.fullmatch(r"[А-Яа-я]+\s+20\d\d", s):  # «Июнь 2025»
        return True
    return False


# ---------------------------------------------------------------------------
# Обязательные праздники
# ---------------------------------------------------------------------------
def _load_holidays(ws) -> list[dict]:
    grid = _build_grid(ws)
    hdr_i = _find_header_row(grid, ["месяц", "дата", "праздник", "тип"])
    if hdr_i is None:
        raise ValueError("не найдена строка заголовка")
    colmap = _column_map(grid[hdr_i])
    out: list[dict] = []
    for i in range(hdr_i + 1, len(grid)):
        row = grid[i]
        name = _clean(_get(colmap, row, "праздник"))
        if not name:
            continue
        out.append({
            "row": i + 1,
            "month": _clean(_get(colmap, row, "месяц")),
            "date_raw": _get(colmap, row, "дата"),
            "name": name,
            "type": _clean(_get(colmap, row, "тип")),
        })
    return out


def _clean(v: Any) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()
