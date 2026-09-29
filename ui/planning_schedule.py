"""Режим «Составить график» (фазы 2a + 2b) — черновой график по правилам и запись.

Раскладывает черновые слоты (дата + проект + тип) по правилам на выбранный
месяц и показывает их на календаре пунктиром (2a). Правила: для каждого проекта
— дни недели, чередование рубрик, дополнительные посты, пропуск праздников.
Кнопка «Записать график в таблицу» (2b) считает план вставки
(checker.schedule_plan) и вставляет новые блоки-посты в лист бренда через
checker.sheets_write — с оформлением по образцу и заголовками новых месяцев.
"""
from __future__ import annotations

import calendar as _cal
import datetime as dt
from typing import Any

import pandas as pd
import streamlit as st
from streamlit_calendar import calendar as st_calendar

from checker import normalize as N
from ui import common as C
from ui.planning_calendar import (PROJECT_COLORS, DEFAULT_COLOR, TYPE_MARKS,
                                   WEEKDAYS_RU, _type_key, _is_holiday, _fmt_date)

# рубрики по умолчанию для выбора (можно выбрать любые из встреченных в таблице)
_DEFAULT_TYPES = ["Информационный", "Развлекательный", "Отгрузка",
                  "Спецпредложение", "Поступление"]

# сколько «дополнительных постов» можно задать на проект
_EXTRA_SLOTS = 2

_SCHEDULE_CSS = """
.fc .ev-draft { outline: 2px dashed #64748B; outline-offset: 1px; }
.fc .fc-daygrid-event { border-radius: 4px; padding: 1px 4px; margin-top: 2px; font-size: 13px; }
.fc .fc-day-sat, .fc .fc-day-sun { background: #F6F7F9; }
.fc .fc-toolbar-title { font-size: 1.4rem; font-weight: 600; }
.fc .fc-toolbar-title::first-letter { text-transform: uppercase; }
.fc .fc-button { background: #1D4ED8; border-color: #1D4ED8; text-transform: none; }
"""


# ---------------------------------------------------------------------------
# Генерация слотов (чистая функция — тестируется отдельно)
# ---------------------------------------------------------------------------
def generate_slots(rules: dict[str, dict], year: int, month: int,
                   holidays: set[dt.date]) -> list[dict]:
    """Разложить черновые слоты по правилам.

    rules[project] = {
        "weekdays": [0..6],        # основные дни — по ним чередуются рубрики
        "types": [str, ...],       # основные рубрики (по кругу)
        "skip_holidays": bool,
        "extra": [{"type": str, "weekdays": [0..6]}, ...],  # доп. посты
    }
    Основной тип назначается по кругу из types в порядке дат. Дополнительные
    посты ставятся своим типом в свои дни, отдельно от основной ротации (в один
    день может выйти и основной, и дополнительный пост). Возвращает список
    {id, date, project, type, kind}.
    """
    slots: list[dict] = []
    ndays = _cal.monthrange(year, month)[1]
    counters = {p: 0 for p in rules}
    for day in range(1, ndays + 1):
        d = dt.date(year, month, day)
        wd = d.weekday()
        for project, r in rules.items():
            if r.get("skip_holidays") and d in holidays:
                continue
            # основной пост — чередование рубрик
            if wd in r.get("weekdays", []):
                types = r.get("types") or [""]
                t = types[counters[project] % len(types)]
                counters[project] += 1
                slots.append({"date": d, "project": project, "type": t,
                              "kind": "main"})
            # дополнительные посты — фиксированный тип в свой день
            for e in r.get("extra", []):
                if wd in e.get("weekdays", []):
                    slots.append({"date": d, "project": project,
                                  "type": e.get("type", ""), "kind": "extra"})
    for i, s in enumerate(slots):
        s["id"] = f"draft|{s['project']}|{s['date'].isoformat()}|{i}"
    return slots


def _holiday_dates(data: "C.AppData") -> set[dt.date]:
    out: set[dt.date] = set()
    for h in getattr(data.wb, "holidays_raw", []) or []:
        d, _ = N.parse_date(h.get("date_raw"))
        if d:
            out.add(d)
    return out


def _draft_events(slots: list[dict], overrides: dict[str, str]) -> list[dict]:
    events = []
    for s in slots:
        d = dt.date.fromisoformat(overrides[s["id"]]) if s["id"] in overrides else s["date"]
        color = PROJECT_COLORS.get(s["project"], DEFAULT_COLOR)
        mark = TYPE_MARKS.get(_type_key(s["type"]), "•")
        events.append({
            "id": s["id"],
            "title": f"✎ {s['project']} · {mark} {s['type']}",
            "start": d.isoformat(),
            "allDay": True,
            "backgroundColor": color,
            "borderColor": color,
            "textColor": "#FFFFFF",
            "classNames": ["ev-draft"],
            "editable": True,
        })
    return events


def _existing_events(df: pd.DataFrame, year: int, month: int) -> list[dict]:
    """Уже существующие посts месяца — приглушённым фоном, без перетаскивания."""
    events = []
    for r in df.itertuples():
        if r.date.year != year or r.date.month != month:
            continue
        color = PROJECT_COLORS.get(r.project, DEFAULT_COLOR)
        events.append({
            "id": f"exist|{r.id}",
            "title": f"{r.project} · {r.type}",
            "start": r.date.isoformat(),
            "allDay": True,
            "backgroundColor": color,
            "borderColor": color,
            "textColor": "#FFFFFF",
            "editable": False,
            "display": "background",
        })
    return events


# ---------------------------------------------------------------------------
# Экран
# ---------------------------------------------------------------------------
def render(data: "C.AppData", df: pd.DataFrame, source_connected: bool) -> None:
    ss = st.session_state
    ss.setdefault("sched_rules", {})
    ss.setdefault("sched_slots", [])
    ss.setdefault("sched_overrides", {})
    ss.setdefault("sched_cal_key", 0)

    brands = list(data.cfg["brands"].keys())
    types_seen = sorted({p.post_type for p in data.posts if p.post_type})
    type_options = list(dict.fromkeys(_DEFAULT_TYPES + types_seen))

    today = C.moscow_today()
    nxt = (today.replace(day=1) + dt.timedelta(days=32)).replace(day=1)

    st.markdown("Задайте правила по проектам, выберите месяц и нажмите «Составить "
                "черновик». Черновые посты появятся на календаре пунктиром — в "
                "таблицу они пока не записываются.")

    # --- выбор месяца ---
    c1, c2 = st.columns(2)
    year = c1.selectbox("Год", [today.year, today.year + 1],
                        index=(0 if nxt.year == today.year else 1), key="sched_year")
    month = c2.selectbox("Месяц", list(range(1, 13)), index=nxt.month - 1,
                         format_func=lambda m: C.MONTHS_NOM[m], key="sched_month")

    # --- правила по проектам ---
    st.subheader("Правила по проектам")
    rules: dict[str, dict] = {}
    for b in brands:
        saved = ss.sched_rules.get(b, {})
        saved_extra = saved.get("extra", [])
        with st.expander(f"{b}", expanded=False):
            st.markdown("**Основные посты** — по этим дням чередуются рубрики")
            wd_labels = st.multiselect(
                "Основные дни недели", WEEKDAYS_RU,
                default=[WEEKDAYS_RU[i] for i in saved.get("weekdays", [])],
                key=f"sched_wd_{b}")
            tsel = st.multiselect(
                "Рубрики (чередуются по кругу)", type_options,
                default=saved.get("types", []), key=f"sched_types_{b}")
            skip = st.checkbox("Пропускать обязательные праздники",
                               value=saved.get("skip_holidays", True),
                               key=f"sched_skip_{b}")

            st.markdown("**Дополнительные посты** — отдельный тип в свой день "
                        "(напр. «Поступление» по пятницам; если основные посты "
                        "только по пятницам — поставьте, например, во вторник)")
            extra = []
            for k in range(_EXTRA_SLOTS):
                sv = saved_extra[k] if k < len(saved_extra) else {}
                ec1, ec2 = st.columns([2, 3])
                et = ec1.selectbox(
                    f"Тип доп. поста {k + 1}", ["—"] + type_options,
                    index=(["—"] + type_options).index(sv["type"])
                    if sv.get("type") in type_options else 0,
                    key=f"sched_extra_type_{b}_{k}")
                ew = ec2.multiselect(
                    f"Дни доп. поста {k + 1}", WEEKDAYS_RU,
                    default=[WEEKDAYS_RU[i] for i in sv.get("weekdays", [])],
                    key=f"sched_extra_wd_{b}_{k}")
                if et != "—" and ew:
                    extra.append({"type": et,
                                  "weekdays": [WEEKDAYS_RU.index(x) for x in ew]})
        weekdays = [WEEKDAYS_RU.index(x) for x in wd_labels]
        if weekdays or extra:
            rules[b] = {"weekdays": weekdays, "types": tsel,
                        "skip_holidays": skip, "extra": extra}

    b1, b2 = st.columns([1, 1])
    if b1.button("Составить черновик", type="primary", use_container_width=True,
                 disabled=not rules):
        ss.sched_rules = rules
        ss.sched_slots = generate_slots(rules, int(year), int(month),
                                        _holiday_dates(data))
        ss.sched_overrides = {}
        ss.sched_cal_key += 1
        st.rerun()
    if b2.button("Очистить черновик", use_container_width=True,
                 disabled=not ss.sched_slots):
        ss.sched_slots = []
        ss.sched_overrides = {}
        ss.sched_cal_key += 1
        st.rerun()

    if not ss.sched_slots:
        if not rules:
            st.info("Пока не заданы правила: выберите хотя бы один день недели у "
                    "проекта.")
        return

    # --- сводка ---
    n = len(ss.sched_slots)
    st.success(f"Черновик: {n} постов на {C.MONTHS_NOM[int(month)].lower()} "
               f"{int(year)}. Двигайте посты на календаре, если нужно поправить дни.")

    # --- календарь черновика ---
    options = {
        "locale": "ru",
        "firstDay": 1,
        "initialView": "dayGridMonth",
        "initialDate": dt.date(int(year), int(month), 1).isoformat(),
        "headerToolbar": {"left": "prev", "center": "title", "right": "next"},
        "height": "auto",
        "fixedWeekCount": False,
        "dayMaxEvents": 5,
        "moreLinkText": "ещё",
        "editable": True,
        "eventDurationEditable": False,
        "eventDisplay": "block",
        "displayEventTime": False,
    }
    events = _existing_events(df, int(year), int(month)) + \
        _draft_events(ss.sched_slots, ss.sched_overrides)
    state = st_calendar(events=events, options=options, custom_css=_SCHEDULE_CSS,
                        callbacks=["eventChange"],
                        key=f"sched_cal_{ss.sched_cal_key}")

    # --- перенос черновых слотов ---
    if state and state.get("callback") == "eventChange":
        ev = state["eventChange"]["event"]
        sid = str(ev["id"])
        new_date = ev["start"][:10]
        stamp = (sid, new_date)
        if sid.startswith("draft|") and ss.get("sched_last_change") != stamp:
            ss.sched_last_change = stamp
            ss.sched_overrides[sid] = new_date
            st.rerun()

    # --- запись в таблицу (шаг 2b) ---
    st.divider()
    if ss.get("sched_apply_done"):
        st.success(ss.pop("sched_apply_done"))
    if ss.get("sched_apply_error"):
        st.error("Не удалось записать график: " + ss.pop("sched_apply_error"))

    confirm = st.checkbox(
        f"Подтверждаю: вставить {n} новых постов в таблицу",
        key="sched_confirm", disabled=not source_connected)
    if st.button("Записать график в таблицу", type="primary",
                 use_container_width=True,
                 disabled=not (source_connected and confirm)):
        _write_schedule(data, ss.sched_slots, ss.sched_overrides)
        st.rerun()
    if not source_connected:
        st.caption("Запись доступна при подключении Google-таблицы через "
                   "сервисный аккаунт с правами «Редактор».")
    else:
        st.caption("Новые строки вставятся в нужный месяц по датам, с оформлением "
                   "по образцу существующих постов. Тексты постов не заполняются.")


def _write_schedule(data: "C.AppData", slots: list[dict],
                    overrides: dict[str, str]) -> None:
    """Вставить черновые слоты в таблицу (по одному листу-бренду за раз).

    Результат кладём в session_state (sched_apply_done / sched_apply_error), а
    при успехе чистим черновик и перечитываем таблицу.
    """
    import io
    from collections import defaultdict

    import openpyxl

    from checker import config_mod, schedule_plan as SP, sheets_write as W

    ss = st.session_state
    sa = C.get_service_account()
    url = (config_mod.load_source().get("url") or "").strip()

    # итоговые даты слотов с учётом переносов, сгруппированные по бренду
    by_brand: dict[str, list[dict]] = defaultdict(list)
    for s in slots:
        d = (dt.date.fromisoformat(overrides[s["id"]])
             if s["id"] in overrides else s["date"])
        by_brand[s["project"]].append({"date": d, "type": s["type"]})

    try:
        wb = openpyxl.load_workbook(io.BytesIO(ss["file_bytes"]))
        written = 0
        for brand, brand_slots in by_brand.items():
            if brand not in wb.sheetnames:
                continue
            layout = SP.parse_layout(wb[brand], brand)
            plan = SP.plan_insertions(layout, brand_slots, wb[brand].max_row)
            if not plan:
                continue
            gid = W.sheet_gid(url, brand, sa)
            tpl = layout.template
            reqs = W.build_insert_requests(
                plan, gid,
                template_row=(tpl.row if tpl else 0),
                template_size=(tpl.size if tpl else 0),
                date_col=layout.date_col, type_col=layout.type_col,
                ncols=layout.ncols)
            W.apply_insertions(url, reqs, sa)
            written += len(plan)
    except Exception as e:  # noqa: BLE001
        ss["sched_apply_error"] = str(e)
        return

    C.fetch_source(force=True)          # перечитать таблицу с новыми постами
    ss["sched_apply_done"] = (f"Записано постов: {written}. Таблица обновлена — "
                              f"новые посты уже видны в календаре.")
    ss["sched_slots"] = []
    ss["sched_overrides"] = {}
    ss["sched_cal_key"] += 1
