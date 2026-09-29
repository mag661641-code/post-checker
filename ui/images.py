"""Экран «Картинки отгрузок».

Сверяет имена картинок в реестре с файлами в папках Google Диска по каждому
бренду (недопустимые имена, транслит, дубли/пропуски номеров, отсутствие
.jpg/.webp, тяжёлые файлы, «сироты», опечатки) и по кнопке проверяет сайт.

Имена читаем из уже подключённой таблицы/Excel, файлы — из Google Диска
(нужен ключ сервисного аккаунта и доступ робота к папкам). Логика —
в checker/shipment_images.py, работа с Диском/сайтом — в checker/shipment_io.py.
"""
from __future__ import annotations

import io
from io import BytesIO
from typing import Any, Optional

import openpyxl
import pandas as pd
import streamlit as st

from checker import config_mod, shipment_images as SI, shipment_io
from checker.dedup import letter_to_index
from checker.models import Level
from ui import common as C


def _sheet_grid(file_bytes: bytes, sheet_name: str) -> Optional[list[list[Any]]]:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    if sheet_name not in wb.sheetnames:
        return None
    ws = wb[sheet_name]
    max_row, max_col = ws.max_row or 0, ws.max_column or 0
    grid = [[ws.cell(r + 1, c + 1).value for c in range(max_col)]
            for r in range(max_row)]
    for mr in ws.merged_cells.ranges:
        top = grid[mr.min_row - 1][mr.min_col - 1]
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                grid[r - 1][c - 1] = top
    return grid


def _col(letter: str) -> Optional[int]:
    letter = (letter or "").strip()
    if not letter:
        return None
    try:
        return letter_to_index(letter)
    except ValueError:
        return None


# ---------------------------------------------------------------------------
# Запуск проверки
# ---------------------------------------------------------------------------
def _run_check(file_bytes: bytes, cfg: dict, sa: Any, with_site: bool,
               progress_cb=None) -> dict:
    max_kb = int(cfg.get("max_kb", 150))
    fuzzy = int(cfg.get("fuzzy_max_diff", 3))
    per_brand = []
    site_queue = []
    skipped = []

    for brand, bcfg in cfg["brands"].items():
        folder_id = (bcfg.get("folder_id") or "").strip()
        if not folder_id:
            skipped.append(brand)
            continue
        grid = _sheet_grid(file_bytes, bcfg["sheet"])
        if grid is None:
            per_brand.append({"brand": brand, "expected_count": 0,
                              "files_count": 0, "problems": [SI.ImgProblem(
                                  brand, Level.ERROR, "Нет листа",
                                  hint="В таблице нет листа «%s»." % bcfg["sheet"])]})
            continue
        name_idx = _col(bcfg.get("name_col"))
        if name_idx is None:
            per_brand.append({"brand": brand, "expected_count": 0,
                              "files_count": 0, "problems": [SI.ImgProblem(
                                  brand, Level.ERROR, "Не задан столбец имени",
                                  hint="Впишите букву столбца в «Настройки»." )]})
            continue
        expected = SI.read_expected_from_grid(
            grid, name_idx, int(bcfg.get("first_data_row", 2)) - 1,
            _col(bcfg.get("jpg_url_col")), _col(bcfg.get("webp_url_col")))
        try:
            drive = shipment_io.load_brand_drive(folder_id, sa)
        except Exception as e:  # noqa: BLE001
            per_brand.append({"brand": brand, "expected_count": len(expected),
                              "files_count": 0, "problems": [SI.ImgProblem(
                                  brand, Level.ERROR, "Не удалось прочитать папку",
                                  hint="Проверьте folder_id и доступ робота к "
                                       "папке. Детали: %s" % e)]})
            continue
        per_brand.append(SI.check_one_brand(brand, expected, drive, bcfg,
                                            max_kb, fuzzy))
        if with_site:
            site_queue += shipment_io.build_site_queue(brand, expected)

    site_problems = []
    if with_site and site_queue:
        site_problems = shipment_io.check_site(site_queue, progress_cb=progress_cb)

    return {"per_brand": per_brand, "site": site_problems, "skipped": skipped}


# ---------------------------------------------------------------------------
# Таблица и Excel
# ---------------------------------------------------------------------------
_COLUMNS = ["Бренд", "Уровень", "Проблема", "Имя в реестре", "Строка",
            "Что на Диске", "Что сделать / подсказка", "Файл"]


def _problems_df(result: dict) -> pd.DataFrame:
    ordered: list[SI.ImgProblem] = []
    for br in result["per_brand"]:
        errs = [p for p in br["problems"] if p.level == Level.ERROR]
        warns = [p for p in br["problems"] if p.level == Level.WARNING]
        if not br["problems"]:
            ordered.append(SI.ImgProblem(br["brand"], None, "Проблем не найдено"))
        else:
            ordered.extend(errs + warns)
    ordered.extend(result.get("site", []))

    rows = []
    for p in ordered:
        rows.append({
            "Бренд": p.brand,
            "Уровень": p.level.title_ru if p.level else "✅",
            "Проблема": p.what,
            "Имя в реестре": p.name,
            "Строка": p.reg_row,
            "Что на Диске": p.disk,
            "Что сделать / подсказка": p.hint,
            "Файл": p.url,
        })
    return pd.DataFrame(rows, columns=_COLUMNS)


def _style_df(df: pd.DataFrame):
    def color_row(row):
        lvl = row.get("Уровень")
        if lvl == Level.ERROR.title_ru:
            return ["background-color: #f8d7da"] * len(row)
        if lvl == Level.WARNING.title_ru:
            return ["background-color: #fff3cd"] * len(row)
        return [""] * len(row)
    return df.style.apply(color_row, axis=1)


def _build_xlsx(result: dict) -> bytes:
    df = _problems_df(result)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Картинки отгрузок", index=False)
        from openpyxl.styles import PatternFill, Font
        ws = writer.book["Картинки отгрузок"]
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FFE8EAED")
        for row in ws.iter_rows(min_row=2):
            val = str(row[1].value)
            color = None
            if val == Level.ERROR.title_ru:
                color = SI.LEVEL_COLOR[Level.ERROR]
            elif val == Level.WARNING.title_ru:
                color = SI.LEVEL_COLOR[Level.WARNING]
            if color:
                fill = PatternFill("solid", fgColor=color)
                for cell in row:
                    cell.fill = fill
    buf.seek(0)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Страница
# ---------------------------------------------------------------------------
def render() -> None:
    file_bytes = st.session_state.get("file_bytes")
    if not file_bytes:
        C.empty_no_file()
        return

    cfg = config_mod.load_images()
    sa = C.get_service_account()

    st.title("Картинки отгрузок")
    st.caption("Сверяет имена картинок в реестре с файлами на Google Диске по "
               "каждому бренду. Диск читается только на чтение.")

    configured = {b: c for b, c in cfg["brands"].items()
                  if (c.get("folder_id") or "").strip()}
    if not sa:
        st.warning("Не добавлен ключ сервисного аккаунта (секрет "
                   "`gcp_service_account_json`). Без него нельзя прочитать папки "
                   "Google Диска. Добавьте ключ и дайте роботу доступ к папкам.")
        return
    if not configured:
        st.info("Не заданы папки Google Диска ни у одного бренда. Откройте "
                "«Настройки» → «🖼 Картинки» и впишите folder_id папок.")
        return

    with_site = st.checkbox(
        "Проверить ещё и сайт (открываются ли картинки). Нужны заполненные "
        "столбцы ссылок в реестре.", value=False)

    if st.button("Проверить картинки", type="primary"):
        try:
            progress = st.progress(0, text="Готовлю проверку…") if with_site \
                else None

            def cb(done, total):
                if progress:
                    progress.progress(int(done / total * 100),
                                      text="Проверяю сайт: %d из %d" % (done, total))

            with st.spinner("Читаю реестры и Google Диск…"):
                result = _run_check(file_bytes, cfg, sa, with_site,
                                    progress_cb=cb if with_site else None)
            if progress:
                progress.empty()
            st.session_state["images_result"] = result
        except Exception as e:  # noqa: BLE001
            st.error("Проверка не выполнена.\n\nЧто случилось: %s" % e)

    result = st.session_state.get("images_result")
    if not result:
        return

    errors = SI.count_errors(result)
    warnings = SI.count_warnings(result)
    m = st.columns(3)
    m[0].metric("Брендов проверено", len(result["per_brand"]))
    m[1].metric("Ошибок", errors)
    m[2].metric("Предупреждений", warnings)

    if result.get("skipped"):
        st.caption("Пропущены (не задан folder_id): "
                   + ", ".join(result["skipped"]))

    df = _problems_df(result)
    brands = ["(все)"] + [br["brand"] for br in result["per_brand"]]
    pick = st.selectbox("Показать бренд", brands)
    show = df if pick == "(все)" else df[df["Бренд"] == pick]
    st.dataframe(
        _style_df(show), use_container_width=True, hide_index=True,
        column_config={"Файл": st.column_config.LinkColumn(
            "Файл", display_text="открыть")})

    st.download_button(
        "⬇️ Скачать Excel (картинки)", data=_build_xlsx(result),
        file_name="proverka_kartinok.xlsx", mime=C.REPORT_MIME)
