"""Экран «Проверка дублей идей».

Читает лист «Идеи» из подключённой Google-таблицы (или загруженного Excel),
ищет повторяющиеся темы по словам и — при включённой смысловой проверке — по
эмбеддингам Gemini. Показывает вердикты цветной таблицей и даёт скачать Excel.

Логика проверки — в checker/dedup.py (без Streamlit), эмбеддинги —
в checker/embeddings.py. Здесь только чтение листа, кеш векторов и отрисовка.
"""
from __future__ import annotations

import hashlib
import io
from io import BytesIO
from typing import Any, Optional

import openpyxl
import pandas as pd
import streamlit as st

from checker import config_mod, dedup, embeddings
from ui import common as C


# ---------------------------------------------------------------------------
# Чтение листа «Идеи»
# ---------------------------------------------------------------------------
def _read_ideas_grid(file_bytes: bytes, sheet_name: str
                     ) -> Optional[list[list[Any]]]:
    """Матрица значений листа «Идеи» (объединённые ячейки развёрнуты сверху)."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    if sheet_name not in wb.sheetnames:
        return None
    ws = wb[sheet_name]
    max_row, max_col = ws.max_row or 0, ws.max_column or 0
    grid = [[ws.cell(r + 1, c + 1).value for c in range(max_col)]
            for r in range(max_row)]
    for mr in ws.merged_cells.ranges:
        top = grid[mr.min_row - 1][mr.min_col - 1]
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                grid[r - 1][c - 1] = top
    return grid


def _thresholds(cfg: dict) -> dedup.Thresholds:
    return dedup.Thresholds(
        dup_limit=float(cfg.get("dup_limit", 0.70)),
        sim_limit=float(cfg.get("sim_limit", 0.50)),
        sem_dup_limit=float(cfg.get("sem_dup_limit", 0.90)),
        sem_sim_limit=float(cfg.get("sem_sim_limit", 0.85)),
        max_idea_length=int(cfg.get("max_idea_length", 400)),
    )


# ---------------------------------------------------------------------------
# Смысловые векторы: кеш по хэшу нормализованного текста (не платим повторно)
# ---------------------------------------------------------------------------
def _vec_key(text: str) -> str:
    return hashlib.md5(dedup.normalize(text).encode("utf-8")).hexdigest()


def _attach_vectors(ideas: list[dedup.Idea], cfg: dict, api_key: str) -> Optional[str]:
    """Проставить idea.vector для пригодных идей. Возвращает текст ошибки или None."""
    cache: dict = st.session_state.setdefault("_dedup_vectors", {})
    usable = [i for i in ideas if i.usable]
    todo = [i for i in usable if _vec_key(i.raw) not in cache]

    if todo:
        progress = st.progress(0.0, text="Считаю смысловые векторы…")

        def cb(done: int, total: int) -> None:
            progress.progress(done / total,
                              text=f"Смысловые векторы: {done} из {total}")

        res = embeddings.embed_texts(
            [i.raw for i in todo], api_key,
            model=cfg.get("embed_model", embeddings.DEFAULT_MODEL),
            dims=int(cfg.get("embed_dims", embeddings.DEFAULT_DIMS)),
            progress_cb=cb)
        progress.empty()
        if not res["ok"]:
            return res["error"]
        for idea, vec in zip(todo, res["vectors"]):
            cache[_vec_key(idea.raw)] = vec

    for idea in usable:
        idea.vector = cache.get(_vec_key(idea.raw))
    return None


# ---------------------------------------------------------------------------
# Таблица и Excel
# ---------------------------------------------------------------------------
_COLUMNS = ["Вердикт", "Строка", "Идея", "Похожесть", "Совпадение (строка)",
            "Совпавшая идея", "Открыть"]


def _verdicts_df(verdicts: list[dedup.Verdict], sheet_name: str) -> pd.DataFrame:
    order = {s: i for i, s in enumerate(dedup.STATUS_ORDER)}
    rows = []
    for v in sorted(verdicts, key=lambda x: (order.get(x.status, 9), -x.score)):
        rows.append({
            "Вердикт": v.label,
            "Строка": v.row,
            "Идея": dedup._trim(v.idea, 120),
            "Похожесть": f"{round(v.score * 100)}% ({v.kind})" if v.score else "",
            "Совпадение (строка)": v.match_row or "",
            "Совпавшая идея": dedup._trim(v.match_text, 120) if v.match_text else "",
            "Открыть": C.sheet_link(v.row, sheet_name) or "",
        })
    return pd.DataFrame(rows, columns=_COLUMNS)


def _build_xlsx(verdicts: list[dedup.Verdict], sheet_name: str) -> bytes:
    df = _verdicts_df(verdicts, sheet_name)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Дубли идей", index=False)
        from openpyxl.styles import PatternFill, Font
        ws = writer.book["Дубли идей"]
        status_by_label = {lbl: s for s, lbl in dedup.STATUS_LABEL.items()}
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FFE8EAED")
        for row in ws.iter_rows(min_row=2):
            status = status_by_label.get(str(row[0].value))
            color = dedup.STATUS_COLOR.get(status)
            if color:
                fill = PatternFill("solid", fgColor=color)
                for cell in row:
                    cell.fill = fill
    buf.seek(0)
    return buf.getvalue()


def _style_df(df: pd.DataFrame):
    status_by_label = {lbl: s for s, lbl in dedup.STATUS_LABEL.items()}
    css = {"taken": "#f4c7c3", "planned": "#fce5cd",
           "similar": "#fff2cc", "free": "#d9ead3", "skipped": "#f3f3f3"}

    def color_row(row):
        status = status_by_label.get(row.get("Вердикт"))
        bg = css.get(status)
        return [f"background-color: {bg}" if bg else ""] * len(row)

    return df.style.apply(color_row, axis=1)


# ---------------------------------------------------------------------------
# Страница
# ---------------------------------------------------------------------------
def render() -> None:
    file_bytes = st.session_state.get("file_bytes")
    if not file_bytes:
        C.empty_no_file()
        return

    cfg = config_mod.load_dedup()
    sheet_name = cfg.get("sheet", "Идеи")
    st.title("Проверка дублей идей")
    st.caption("Ищет повторяющиеся темы на листе «%s»: по словам всегда, по "
               "смыслу — если включить и добавлен ключ Google AI." % sheet_name)

    grid = _read_ideas_grid(file_bytes, sheet_name)
    if grid is None:
        st.warning(f"В таблице нет листа «{sheet_name}». Проверьте название на "
                   f"странице «Настройки» → «Дубли идей» (или переименуйте лист).")
        return

    th = _thresholds(cfg)
    ideas = dedup.read_ideas_from_grid(
        grid, int(cfg["col_idea"]), int(cfg["col_published"]),
        int(cfg["row_start"]), th)
    usable = [i for i in ideas if i.usable]
    if not usable:
        st.info("На листе не нашлось идей для проверки. Проверьте настройки "
                "столбцов на странице «Настройки».")
        return

    api_key = C.get_ai_key()
    can_semantic = bool(cfg.get("use_semantic", True) and api_key)

    c1, c2 = st.columns([3, 1])
    with c1:
        st.write(f"Идей к проверке: **{len(usable)}** "
                 f"(из {len(ideas)} строк с текстом).")
    with c2:
        semantic = st.toggle("Проверка по смыслу", value=can_semantic,
                             disabled=not api_key,
                             help=("Косинус эмбеддингов Gemini. Нужен ключ "
                                   "Google AI (см. Настройки → Нейросеть).")
                             if not api_key else None)

    if not api_key:
        st.caption("🔵 Смысловая проверка недоступна: не добавлен ключ Google AI. "
                   "Работает проверка по словам.")

    if st.button("Проверить дубли", type="primary"):
        if semantic:
            err = _attach_vectors(ideas, cfg, api_key)
            if err:
                st.error("Смысловая проверка не выполнена: %s\n\n"
                         "Показываю результат по словам." % err)
        with st.spinner("Сравниваю идеи…"):
            verdicts = dedup.run_dedup(ideas, th)
        st.session_state["dedup_result"] = {
            "verdicts": verdicts, "sheet": sheet_name}

    result = st.session_state.get("dedup_result")
    if not result:
        return

    verdicts = result["verdicts"]
    counts = dedup.summarize(verdicts)
    m = st.columns(5)
    m[0].metric("🔴 Занято", counts["taken"])
    m[1].metric("🟠 В плане", counts["planned"])
    m[2].metric("🟡 Похоже", counts["similar"])
    m[3].metric("🟢 Свободно", counts["free"])
    m[4].metric("⚪ Пропущено", counts["skipped"])

    only_problems = st.toggle("Показывать только совпадения", value=True,
                              key="dedup_only")
    shown = [v for v in verdicts if v.status in ("taken", "planned", "similar")] \
        if only_problems else verdicts
    if not shown:
        st.success("Дублей не найдено — все идеи уникальны.")
    else:
        df = _verdicts_df(shown, result["sheet"])
        st.dataframe(
            _style_df(df), use_container_width=True, hide_index=True,
            column_config={"Открыть": st.column_config.LinkColumn(
                "Открыть", display_text="в таблице")})

    st.download_button(
        "⬇️ Скачать Excel (дубли идей)",
        data=_build_xlsx(verdicts, result["sheet"]),
        file_name="proverka_dublej_idej.xlsx",
        mime=C.REPORT_MIME)

    with st.expander("Калибровка порогов по смыслу (для похожих тем)"):
        _calibration(usable)


def _calibration(usable: list[dedup.Idea]) -> None:
    """Распределение смысловой похожести по реальным данным — чтобы пороги не
    угадывать (перенос функции calibrate из исходного скрипта)."""
    with_vec = [i for i in usable if i.vector is not None]
    if len(with_vec) < 5:
        st.caption("Нужно ≥5 идей с посчитанными векторами. Запустите проверку "
                   "с включённой смысловой проверкой.")
        return
    scores = []
    for a in range(len(with_vec)):
        for b in range(a + 1, len(with_vec)):
            scores.append(dedup.cosine(with_vec[a].vector, with_vec[b].vector))
    scores.sort(reverse=True)

    def pct(p: float) -> int:
        return round(scores[min(len(scores) - 1, int(len(scores) * p))] * 100)

    st.write(f"Пар идей: {len(scores)}. Медиана (фон тематики): **{pct(0.5)}%**.")
    st.write(f"Рекомендация: SEM_DUP ≈ **{pct(0.002)}%**, "
             f"SEM_SIM ≈ **{pct(0.01)}%**. Ниже медианы опускаться бессмысленно.")
    st.write("Самые похожие пары: " +
             ", ".join(f"{round(s * 100)}%" for s in scores[:8]))
