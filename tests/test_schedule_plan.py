"""Расчёт плана вставки постов графика в лист бренда (без записи)."""
import datetime as dt
from io import BytesIO

import openpyxl

from checker import schedule_plan as SP


def _sheet_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "МПЭ"
    ws.append(["Дата", "Соцсеть", "Бренд", "Ссылка", "Формат", "Тип", "Пост", "Фото"])
    # заголовок месяца на всю ширину
    ws.append(["Сентябрь 2026", None, None, None, None, None, None, None])
    ws.merge_cells("A2:H2")
    # пост 02.09 — блок 2 строки
    ws.append([dt.date(2026, 9, 2), "Telegram", "МПЭ", "", "Пост", "Информационный", "текст", ""])
    ws.append([None, "Вконтакте", None, "", "Пост", None, None, ""])
    ws.merge_cells("A3:A4")
    ws.merge_cells("F3:F4")
    # пост 20.09 — блок 2 строки
    ws.append([dt.date(2026, 9, 20), "Telegram", "МПЭ", "", "Пост", "Отгрузка", "текст", ""])
    ws.append([None, "Вконтакте", None, "", "Пост", None, None, ""])
    ws.merge_cells("A5:A6")
    ws.merge_cells("F5:F6")
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return openpyxl.load_workbook(buf)


def test_layout_parsed():
    ws = _sheet_bytes()["МПЭ"]
    lay = SP.parse_layout(ws, "МПЭ")
    assert lay.date_col == 0 and lay.type_col == 5
    assert lay.block_size == 2
    assert len(lay.sections) == 1 and (lay.sections[0].year, lay.sections[0].month) == (2026, 9)
    assert {b.date for b in lay.posts} == {dt.date(2026, 9, 2), dt.date(2026, 9, 20)}


def test_insert_in_date_order_existing_section():
    wb = _sheet_bytes()
    ws = wb["МПЭ"]
    lay = SP.parse_layout(ws, "МПЭ")
    # новый пост 10.09 должен встать ПЕРЕД постом 20.09 (строка 5)
    plan = SP.plan_insertions(lay, [{"date": dt.date(2026, 9, 10), "type": "Отгрузка"}], ws.max_row)
    assert len(plan) == 1
    assert plan[0].at_row == 5
    assert plan[0].new_header is None
    assert plan[0].size == 2


def test_new_month_creates_header_at_end():
    wb = _sheet_bytes()
    ws = wb["МПЭ"]
    lay = SP.parse_layout(ws, "МПЭ")
    plan = SP.plan_insertions(lay, [{"date": dt.date(2026, 10, 5), "type": "Поступление"}], ws.max_row)
    assert plan[0].new_header == "Октябрь 2026"
    assert plan[0].at_row == ws.max_row + 1


def test_insert_at_section_end_when_latest():
    wb = _sheet_bytes()
    ws = wb["МПЭ"]
    lay = SP.parse_layout(ws, "МПЭ")
    # 25.09 — позже всех в сентябре → в конец раздела (после строки 6)
    plan = SP.plan_insertions(lay, [{"date": dt.date(2026, 9, 25), "type": "Отгрузка"}], ws.max_row)
    assert plan[0].at_row == 7  # hi(6) + 1
    assert plan[0].new_header is None
